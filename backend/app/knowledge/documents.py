"""Document readers: PDF, DOCX, TXT, PPTX and optional OCR for photos."""
from pathlib import Path


def read_file(path: str) -> str:
    """Best-effort text extraction by extension."""
    ext = Path(path).suffix.lower()
    if ext == ".pdf":
        return _read_pdf(path)
    if ext in (".docx", ".doc"):
        return _read_docx(path)
    if ext in (".pptx", ".ppt"):
        return _read_pptx(path)
    if ext in (".txt", ".md", ".csv"):
        return Path(path).read_text(encoding="utf-8", errors="ignore")
    if ext in (".png", ".jpg", ".jpeg", ".webp", ".bmp"):
        return _read_image(path)
    return ""


def _read_pdf(path: str) -> str:
    from pypdf import PdfReader

    reader = PdfReader(path)
    return "\n\n".join(p.extract_text() or "" for p in reader.pages)


def _read_docx(path: str) -> str:
    import docx2txt

    return docx2txt.process(path) or ""


def _read_pptx(path: str) -> str:
    from pptx import Presentation

    prs = Presentation(path)
    parts = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                parts.append(shape.text_frame.text)
    return "\n".join(p for p in parts if p)


def _read_image(path: str) -> str:
    try:
        from PIL import Image
        import pytesseract

        img = Image.open(path)
        return pytesseract.image_to_string(img) or ""
    except Exception:
        return ""


def file_handled_extension(ext: str) -> bool:
    return ext.lower() in (".pdf", ".docx", ".doc", ".pptx", ".ppt", ".txt", ".md",
                           ".png", ".jpg", ".jpeg", ".webp", ".bmp", ".csv")